#!/usr/bin/env python3
"""総合エネルギー統計「時系列表」の Excel から public/data の cube を書く。設計は docs/data-sources.md。"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data/raw/stte_jikeiretu2024fykaku.xlsx"
OUT = ROOT / "public/data"

LABEL_COL = 2
UNIT_COL = 38


@dataclass
class Item:
    code: str
    label: str
    parent: str | None = None


@dataclass
class View:
    name: str
    sheet: str
    unit: str
    items: list[Item]
    groups: list[tuple[str, str]] = field(default_factory=list)


VIEWS = [
    View(
        "supply",
        "1.一次エネルギー国内供給",
        "PJ",
        [
            Item("coal", "石炭"),
            Item("oil", "石油"),
            Item("gas", "天然ガス・都市ガス"),
            Item("renewable", "再生可能（水力除く）"),
            Item("hydro", "水力（揚水除く）"),
            Item("unused", "未活用"),
            Item("nuclear", "原子力"),
        ],
    ),
    View(
        "final",
        "2.最終消費(エネルギー源別)",
        "PJ",
        [
            Item("coal", "石炭"),
            Item("oil", "石油"),
            Item("gas", "天然ガス・都市ガス"),
            Item("renewable", "再生可能"),
            Item("unused", "未活用"),
            Item("electricity", "電力"),
            Item("heat", "熱"),
        ],
    ),
    View(
        "sector",
        "3.最終消費(部門別)",
        "PJ",
        [
            Item("agriculture", "農林水産鉱業建設", "business"),
            Item("manufacturing", "製造業", "business"),
            Item("commercial", "業務他", "business"),
            Item("household", "家庭", "household"),
            Item("passenger", "旅客", "transport"),
            Item("freight", "貨物", "transport"),
        ],
        [("business", "企業・事業所他"), ("household", "家庭"), ("transport", "運輸")],
    ),
    View(
        "power",
        "4.電源構成(発電量）",
        "億kWh",
        [
            Item("nuclear", "原子力"),
            Item("coal", "石炭"),
            Item("gas", "天然ガス"),
            Item("oil", "石油等"),
            Item("hydro", "水力"),
            Item("solar", "太陽光"),
            Item("wind", "風力"),
            Item("geothermal", "地熱"),
            Item("biomass", "バイオマス"),
        ],
    ),
]


def label(ws: Worksheet, row: int) -> str:
    v = ws.cell(row, LABEL_COL).value
    return "" if v is None else str(v).strip().replace("\u3000", "")


def blocks(ws: Worksheet) -> dict[str, tuple[int, dict[int, int]]]:
    """単位 → (見出し行, {年度: 列})。見出しは B 列が「年度」の行で、単位はその1行上の AL 列。"""
    out: dict[str, tuple[int, dict[int, int]]] = {}
    for row in range(1, ws.max_row + 1):
        if label(ws, row) != "年度":
            continue
        unit = str(ws.cell(row - 1, UNIT_COL).value or "").strip("()（）")
        years = {
            int(c.value): c.column
            for c in ws[row]
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
        }
        out[unit] = (row, years)
    return out


def read_block(ws: Worksheet, header: int, labels: list[str]) -> dict[str, int]:
    """見出し行の下から空行までを読み、ラベル → 行。"""
    rows: dict[str, int] = {}
    row = header + 1
    while label(ws, row) != "":
        rows[label(ws, row)] = row
        row += 1
    missing = [x for x in labels if x not in rows]
    if missing:
        raise SystemExit(f"{ws.title}: 項目が見つからない {missing}")
    return rows


def num(v: object) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def build(wb: openpyxl.Workbook, view: View) -> dict:
    ws = wb[view.sheet]
    found = blocks(ws)
    header, cols = found[view.unit]
    labels = [i.label for i in view.items]
    rows = read_block(ws, header, labels + ["合計"])

    years = sorted(
        y for y, c in cols.items() if all(num(ws.cell(rows[x], c).value) is not None for x in labels)
    )
    if years != list(range(years[0], years[-1] + 1)):
        raise SystemExit(f"{view.sheet}: 年度が連続していない {years}")

    values = {i.code: [num(ws.cell(rows[i.label], cols[y]).value) for y in years] for i in view.items}
    totals = [num(ws.cell(rows["合計"], cols[y]).value) for y in years]

    share_header, share_cols = found["構成比"]
    share_rows = read_block(ws, share_header, labels)
    for k, y in enumerate(years):
        total = sum(values[i.code][k] for i in view.items)
        if abs(total - totals[k]) > 1e-6 * totals[k]:
            raise SystemExit(f"{view.sheet} {y}: 内訳の和 {total} と合計 {totals[k]} が合わない")
        for i in view.items:
            official = num(ws.cell(share_rows[i.label], share_cols[y]).value)
            if official is None or abs(values[i.code][k] / total - official) > 1e-6:
                raise SystemExit(f"{view.sheet} {y} {i.label}: 構成比 {official} と合わない")

    result: dict = {
        "unit": view.unit,
        "items": [
            {"code": i.code, "label": i.label, "level": 1, **({"parent": i.parent} if i.parent else {})}
            for i in view.items
        ],
    }
    if view.groups:
        result["groups"] = [{"code": c, "label": l, "level": 0} for c, l in view.groups]
    result["dims"] = [
        {"name": "item", "codes": [i.code for i in view.items]},
        {"name": "year", "codes": [str(y) for y in years]},
    ]
    result["measures"] = {"value": [round(v, 3) for i in view.items for v in values[i.code]]}
    return result


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"{XLSX.relative_to(ROOT)} が無い。npm run fetch を先に実行する。")
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    OUT.mkdir(parents=True, exist_ok=True)
    for view in VIEWS:
        data = build(wb, view)
        years = data["dims"][1]["codes"]
        (OUT / f"{view.name}.json").write_text(
            json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8"
        )
        print(f"{view.name}.json  {len(view.items)} items × {years[0]}–{years[-1]}")


if __name__ == "__main__":
    main()
